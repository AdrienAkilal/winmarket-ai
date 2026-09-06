from pathlib import Path
from datetime import datetime
from typing import Optional
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import cm
from src.core.config import OUTPUT_DIR, LLM_TEMPERATURE_GENERATION
from src.core.models import AOContext, ScoringResult


def _safe_name(titre: str) -> str:
    return "".join(c if c.isalnum() else "_" for c in titre[:40]) or "ao"


class DocumentGenerator:
    def __init__(self, output_dir: Path = OUTPUT_DIR):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    _DOC_SYSTEM = (
        "Tu es rédacteur avant-vente senior dans une ESN française spécialisée en transformation digitale, "
        "développement applicatif, data/IA, cloud et cybersécurité. Tu rédiges des dossiers de candidature "
        "qui remportent des marchés. Ton écriture est précise, dense, orientée valeur client. "
        "Tu maîtrises les codes de la réponse aux appels d'offres : ton institutionnel mais engageant, "
        "arguments factuels et vérifiables, focus sur les bénéfices client plutôt que sur tes propres "
        "capacités. Tu adaptes TOUJOURS ton discours au secteur, à la culture et aux enjeux spécifiques "
        "du client. Tu évites absolument : le remplissage, les banalités, les phrases sans substance, "
        "les formules passe-partout ('fort de notre expérience...', 'notre approche éprouvée...', "
        "'nous mettons tout en oeuvre...'). Chaque phrase doit apporter une information ou un argument "
        "concret qui aide le client à choisir notre ESN."
    )

    def _build_context(self, ao: AOContext, result: ScoringResult) -> str:
        refs_text = "\n".join([f"  • {ev.source} — pertinence {min(int(ev.score*400),100)}%" for ev in result.evidence_pack[:4]]) or "  • Aucune référence disponible"
        return f"""APPEL D'OFFRES
Titre     : {ao.titre}
Client    : {ao.client}
Secteur   : {ao.secteur or 'Non précisé'}
Budget    : {f"{ao.budget_estime:,.0f} €" if ao.budget_estime else 'Non communiqué'}
Durée     : {f"{ao.duree_projet_mois} mois" if ao.duree_projet_mois else 'Non précisée'}
Techs     : {', '.join(ao.technologies_demandees) or 'Non précisées'}
Livrables : {', '.join(ao.livrables[:4]) or 'Non précisés'}
Contraintes: {' | '.join(ao.contraintes[:3]) or 'Aucune'}
Certif.   : {', '.join(ao.certifications_obligatoires) or 'Aucune'}
Deadline  : {ao.deadline_reponse or 'Non précisée'}

SCORING
Décision  : {result.decision} ({result.score_global}/100)
Bloquants : {' | '.join(result.criteres_bloquants) or 'Aucun'}
Forces    : {' | '.join(result.forces[:3]) or 'Aucune'}
Faiblesses: {' | '.join(result.faiblesses[:3]) or 'Aucune'}
Références mobilisables :
{refs_text}"""

    def _generate_ai_content(self, ao: AOContext, result: ScoringResult, llm) -> dict:
        """Génère le contenu IA adapté à la décision (GO / GO SOUS RÉSERVE / NO-GO)."""
        if not llm.enabled:
            return {}

        ctx = self._build_context(ao, result)
        decision = result.decision

        if decision == "NO-GO":
            prompt = f"""Tu dois rédiger un mémo de non-réponse professionnel pour cet appel d'offres.

{ctx}

MISSION : Rédige les sections du mémo de non-réponse. Sois direct, professionnel, factuel.

**"note_refus"** — 3 phrases. Synthèse claire et factuelle de la décision de ne pas répondre, ancrée dans les critères bloquants identifiés. Ton : professionnel, pas défensif.

**"analyse_blocages"** — 4 à 5 phrases. Explication détaillée de chaque point bloquant avec son impact concret sur notre capacité à répondre ou livrer. Citer les certifications manquantes, technologies hors périmètre, ou risques contractuels spécifiques à CET AO.

**"risques_si_reponse"** — 3 phrases. Ce qui se passerait si on répondait malgré les blocages : risques opérationnels, financiers, réputationnels. Être factuel, pas dramatique.

**"conditions_reconsideration"** — Liste de 2 à 4 items (strings). Actions concrètes et mesurables qui nous permettraient de répondre à un AO similaire à l'avenir. Chaque item doit être spécifique et actionnable (ex: "Obtenir la certification X d'ici 6 mois via...").

**"alternatives_proposees"** — 2 à 3 phrases. Ce qu'on peut proposer à {ao.client} à la place : périmètre réduit, autre type de mission, partenariat avec un tiers qualifié. Maintenir la relation commerciale.

Réponds UNIQUEMENT en JSON valide. "conditions_reconsideration" est une liste de strings, les autres sont des strings.
"""
        elif "RESERVE" in decision:
            prompt = f"""Tu dois rédiger un dossier de candidature conditionnel pour cet appel d'offres.
La décision est GO SOUS RÉSERVE : on peut répondre mais des points doivent être levés avant soumission.

{ctx}

MISSION : Rédige les sections du dossier conditionnel. Intègre les réserves comme une donnée structurante, pas comme un problème caché.

**"resume_executif"** — 3 phrases. Montrer qu'on maîtrise le projet tout en signalant notre transparence sur les points à valider. Ton : confiant mais honnête.

**"comprehension_besoin"** — 4 phrases. Démontrer la compréhension profonde du besoin de {ao.client} : enjeux métier {ao.secteur or ''}, défis techniques, facteurs clés de succès de CE projet.

**"reserves_et_conditions"** — 4 à 5 phrases. Section centrale : exposer clairement les réserves, leur nature, et le plan concret pour les lever avant la date de soumission ({ao.deadline_reponse or 'la date limite'}). Ton : transparent et professionnel.

**"plan_action_go"** — Liste de 3 à 4 items (strings). Actions concrètes à mener avant la soumission pour transformer ce GO SOUS RÉSERVE en GO définitif. Format : "Action : [action] — Responsable : [qui] — Délai : [quand]".

**"methodologie"** — 4 phrases. Approche projet spécifique à CE projet : phases liées aux livrables ({', '.join(ao.livrables[:3]) or 'non précisés'}), gestion des risques identifiés, jalons de validation.

**"equipe_proposee"** — 3 phrases. Profils clés pour la stack {', '.join(ao.technologies_demandees[:3]) or 'demandée'} et le secteur {ao.secteur or 'du client'}.

**"valeur_ajoutee"** — Liste de 3 items (strings). Raisons objectives de nous choisir malgré les réserves. Chaque item doit être vérifiable et spécifique à CET AO.

**"conclusion"** — 2 phrases. Engagement conditionnel clair, invitation à un échange avant soumission pour valider les points de réserve.

Réponds UNIQUEMENT en JSON valide. "plan_action_go" et "valeur_ajoutee" sont des listes de strings, les autres sont des strings.
"""
        else:  # GO
            prompt = f"""Tu dois rédiger un dossier de candidature complet et percutant pour cet appel d'offres.
La décision est GO : on répond avec ambition.

{ctx}

MISSION : Rédige les sections du dossier. Chaque phrase doit apporter un argument concret qui aide {ao.client} à nous choisir.

**"resume_executif"** — 3 à 4 phrases. (1) Reformulation de l'enjeu stratégique de {ao.client}, (2) Notre positionnement différenciant sur ce type de projet, (3) Ce qu'on va livrer concrètement.
ÉVITER : Commencer par "Notre ESN...", les superlatifs non étayés.

**"comprehension_besoin"** — 4 à 5 phrases. Démontrer la compréhension profonde : enjeux métier {ao.secteur or ''}, besoins fonctionnels précis, enjeux techniques ({', '.join(ao.technologies_demandees[:3]) or 'stack demandée'}), facteurs clés de succès.

**"methodologie"** — 5 phrases. Approche spécifique à CE projet : phases liées aux livrables ({', '.join(ao.livrables[:3]) or 'non précisés'}), gestion des risques, jalons de validation, gouvernance adaptée au secteur.
ÉVITER : "Approche Agile Scrum" sans justification, phases génériques sans lien avec les livrables.

**"equipe_proposee"** — 4 phrases. Chef de projet (profil sectoriel), experts techniques sur la stack demandée, expertises spécialisées (data/sécurité/UX selon l'AO), organisation et continuité.

**"valeur_ajoutee"** — Liste de 3 à 4 items (strings). Raisons objectives et vérifiables de nous choisir. Pas de formules génériques — chaque item doit être spécifique à CET AO.

**"conclusion"** — 2 à 3 phrases. Réaffirmation des enjeux de {ao.client}, proposition d'échange pour approfondir, engagement sobre et professionnel.

Réponds UNIQUEMENT en JSON valide. "valeur_ajoutee" est une liste de strings, les autres sont des strings.
"""

        try:
            return llm.json_complete(prompt, system=self._DOC_SYSTEM, temperature=LLM_TEMPERATURE_GENERATION, max_tokens=3000) or {}
        except Exception:
            return {}

    def generate_docx(self, ao: AOContext, result: ScoringResult, llm=None) -> Path:
        safe = _safe_name(ao.titre)
        ts = datetime.now().strftime("%Y%m%d_%H%M")
        path = self.output_dir / f"candidature_{safe}_{ts}.docx"

        ai = result.ai_content or (self._generate_ai_content(ao, result, llm) if llm else {})
        decision = result.decision

        doc = Document()

        # ── Header commun ──
        if decision == "NO-GO":
            doc.add_heading("Mémo de non-réponse — WinMarket AI", 0)
        elif "RESERVE" in decision:
            doc.add_heading("Dossier de candidature conditionnel — WinMarket AI", 0)
        else:
            doc.add_heading("Dossier de candidature — WinMarket AI", 0)

        p = doc.add_paragraph()
        p.add_run("Décision : ").bold = True
        p.add_run(result.decision + "   ")
        p.add_run("Score : ").bold = True
        p.add_run(f"{result.score_global}/100   |   ")
        p.add_run("Client : ").bold = True
        p.add_run(f"{ao.client}   |   {ao.titre[:80]}")

        # ════════════════════════════════
        # CAS NO-GO
        # ════════════════════════════════
        if decision == "NO-GO":
            doc.add_heading("1. Décision et synthèse", 1)
            doc.add_paragraph(ai.get("note_refus") or
                f"Après analyse, notre ESN a décidé de ne pas répondre à l'appel d'offres "
                f"'{ao.titre}' de {ao.client}. Les critères bloquants identifiés rendent "
                f"une réponse dans les conditions actuelles trop risquée.")

            if result.criteres_bloquants:
                doc.add_heading("2. Critères bloquants", 1)
                doc.add_paragraph(ai.get("analyse_blocages") or "")
                for b in result.criteres_bloquants:
                    doc.add_paragraph(b, style="List Bullet")

            doc.add_heading("3. Risques si réponse quand même", 1)
            doc.add_paragraph(ai.get("risques_si_reponse") or
                "Répondre malgré les blocages exposerait notre ESN à des risques opérationnels "
                "et réputationnels significatifs.")

            doc.add_heading("4. Conditions pour reconsidérer", 1)
            conditions = ai.get("conditions_reconsideration") or result.recommandations
            items = conditions if isinstance(conditions, list) else [conditions]
            for item in items:
                doc.add_paragraph(str(item), style="List Bullet")

            doc.add_heading("5. Alternatives et maintien de la relation", 1)
            doc.add_paragraph(ai.get("alternatives_proposees") or
                f"Nous recommandons de maintenir la relation avec {ao.client} et de "
                f"nous positionner sur les prochains marchés une fois les blocages levés.")

            doc.add_heading("6. Scoring détaillé (référence interne)", 1)
            for c in result.criteres:
                p = doc.add_paragraph()
                p.add_run(f"{c.nom} ({int(c.poids)}%) : ").bold = True
                p.add_run(f"{c.score:.0f}/100 — {c.justification}")

        # ════════════════════════════════
        # CAS GO SOUS RÉSERVE
        # ════════════════════════════════
        elif "RESERVE" in decision:
            doc.add_paragraph(
                "⚠️  DOCUMENT CONDITIONNEL — Ce dossier est préparé sous réserve de validation "
                "des points listés en section 3. Ne pas soumettre avant levée des réserves.",
                style="Intense Quote" if "Intense Quote" in [s.name for s in doc.styles] else "Normal"
            )

            doc.add_heading("1. Résumé exécutif", 1)
            doc.add_paragraph(ai.get("resume_executif") or
                f"Cette candidature pour le projet '{ao.titre}' de {ao.client} est préparée "
                f"sous réserve de validation de points spécifiques détaillés ci-après.")

            doc.add_heading("2. Compréhension du besoin", 1)
            doc.add_paragraph(ai.get("comprehension_besoin") or
                f"Le projet de {ao.client} porte sur {', '.join(ao.technologies_demandees[:3]) or 'les technologies demandées'}.")
            if ao.contraintes:
                for c in ao.contraintes[:4]:
                    doc.add_paragraph(c, style="List Bullet")

            doc.add_heading("3. Réserves et conditions à lever", 1)
            doc.add_paragraph(ai.get("reserves_et_conditions") or
                "Les points suivants doivent être validés avant soumission définitive :")
            if result.criteres_bloquants:
                for b in result.criteres_bloquants:
                    doc.add_paragraph(f"🔴 {b}", style="List Bullet")
            faiblesses_notables = [c for c in result.criteres if 55 <= c.score < 70]
            for f in faiblesses_notables[:3]:
                doc.add_paragraph(f"🟡 {f.nom} : score {f.score:.0f}/100 — {f.justification}", style="List Bullet")

            doc.add_heading("4. Plan d'action pour GO définitif", 1)
            plan = ai.get("plan_action_go") or result.recommandations
            items = plan if isinstance(plan, list) else [plan]
            for item in items:
                doc.add_paragraph(str(item), style="List Bullet")

            doc.add_heading("5. Méthodologie proposée", 1)
            doc.add_paragraph(ai.get("methodologie") or
                "Approche par phases liées aux livrables attendus, avec points de validation client à chaque jalon.")

            doc.add_heading("6. Équipe projet", 1)
            doc.add_paragraph(ai.get("equipe_proposee") or
                f"Équipe constituée autour des compétences clés : {', '.join(ao.technologies_demandees[:3]) or 'demandées'}.")

            doc.add_heading("7. Références similaires", 1)
            if result.rag_synthesis:
                doc.add_paragraph(result.rag_synthesis)
            for ev in result.evidence_pack[:3]:
                p = doc.add_paragraph()
                p.add_run(f"{ev.source}  |  Pertinence : {min(int(ev.score*400),100)}%").bold = True
                doc.add_paragraph(ev.content[:400])

            valeur = ai.get("valeur_ajoutee")
            if valeur:
                doc.add_heading("8. Notre valeur ajoutée", 1)
                items = valeur if isinstance(valeur, list) else [valeur]
                for item in items:
                    doc.add_paragraph(str(item), style="List Bullet")

            doc.add_heading("9. Conclusion", 1)
            doc.add_paragraph(ai.get("conclusion") or
                f"Sous réserve de validation des points mentionnés, nous sommes prêts à nous engager "
                f"pleinement sur ce projet. Nous proposons un échange avec {ao.client} avant soumission.")

        # ════════════════════════════════
        # CAS GO
        # ════════════════════════════════
        else:
            doc.add_heading("1. Résumé exécutif", 1)
            doc.add_paragraph(ai.get("resume_executif") or
                f"Notre ESN répond à l'appel d'offres '{ao.titre}' émis par {ao.client} avec "
                f"un positionnement fort sur la stack demandée et des références sectorielles probantes.")

            doc.add_heading("2. Compréhension du besoin", 1)
            doc.add_paragraph(ai.get("comprehension_besoin") or
                f"Le projet de {ao.client} porte sur {', '.join(ao.technologies_demandees[:3]) or 'les technologies demandées'}.")
            if ao.contraintes:
                doc.add_paragraph("Contraintes identifiées :")
                for c in ao.contraintes[:5]:
                    doc.add_paragraph(c, style="List Bullet")

            doc.add_heading("3. Évaluation de l'adéquation", 1)
            forces_criteres = [c for c in result.criteres if c.score >= 78]
            for c in forces_criteres[:5]:
                p = doc.add_paragraph()
                p.add_run(f"✓ {c.nom} : ").bold = True
                p.add_run(f"{c.score:.0f}/100 — {c.justification}")

            doc.add_heading("4. Méthodologie proposée", 1)
            doc.add_paragraph(ai.get("methodologie") or
                "Approche par phases liées aux livrables attendus, avec points de validation client à chaque jalon clé.")

            doc.add_heading("5. Équipe projet", 1)
            doc.add_paragraph(ai.get("equipe_proposee") or
                f"Équipe dédiée avec expertise sur {', '.join(ao.technologies_demandees[:3]) or 'les technologies demandées'}.")

            doc.add_heading("6. Références et expériences similaires", 1)
            if result.rag_synthesis:
                doc.add_paragraph(result.rag_synthesis)
            for ev in result.evidence_pack[:4]:
                p = doc.add_paragraph()
                p.add_run(f"{ev.source}  |  Pertinence : {min(int(ev.score*400),100)}%").bold = True
                doc.add_paragraph(ev.content[:500])

            valeur = ai.get("valeur_ajoutee")
            if valeur:
                doc.add_heading("7. Notre valeur ajoutée", 1)
                items = valeur if isinstance(valeur, list) else [valeur]
                for item in items:
                    doc.add_paragraph(str(item), style="List Bullet")

            doc.add_heading("8. Recommandations avant-vente", 1)
            for r in result.recommandations:
                doc.add_paragraph(r, style="List Bullet")

            doc.add_heading("9. Conclusion", 1)
            doc.add_paragraph(ai.get("conclusion") or
                f"Notre ESN est pleinement mobilisée pour répondre à ce projet de {ao.client}. "
                f"Nous sommes disponibles pour tout échange complémentaire.")

        doc.save(path)
        return path

    def generate_pdf(self, ao: AOContext, result: ScoringResult, llm=None) -> Path:
        safe = _safe_name(ao.titre)
        ts = datetime.now().strftime("%Y%m%d_%H%M")
        path = self.output_dir / f"rapport_decision_{safe}_{ts}.pdf"

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle("CustomTitle", parent=styles["Title"], fontSize=18, spaceAfter=12)
        h2_style = ParagraphStyle("CustomH2", parent=styles["Heading2"], fontSize=13, spaceBefore=10)
        normal = styles["Normal"]

        # Decision color
        decision_color = colors.green if result.decision == "GO" else (
            colors.orange if "RESERVE" in result.decision else colors.red
        )
        decision_style = ParagraphStyle(
            "Decision", parent=styles["Heading1"],
            textColor=decision_color, fontSize=16
        )

        ai = result.ai_content or (self._generate_ai_content(ao, result, llm) if llm else {})

        story = [
            Paragraph("Rapport de décision — WinMarket AI", title_style),
            HRFlowable(width="100%", thickness=1, color=colors.grey),
            Spacer(1, 12),
            Paragraph(f"Appel d'offres : {ao.titre}", h2_style),
            Paragraph(f"Client : {ao.client}", normal),
            Paragraph(f"Secteur : {ao.secteur or 'Non renseigné'}", normal),
            Spacer(1, 8),
            Paragraph(f"DÉCISION : {result.decision}", decision_style),
            Paragraph(f"Score global : {result.score_global} / 100", h2_style),
            Spacer(1, 12),
        ]

        # Résumé exécutif IA
        resume = ai.get("resume_executif")
        if resume:
            story.append(Paragraph("Résumé exécutif", h2_style))
            story.append(Paragraph(resume, normal))
            story.append(Spacer(1, 8))

        # Synthèse RAG
        if result.rag_synthesis:
            story.append(Paragraph("Analyse des références internes", h2_style))
            story.append(Paragraph(result.rag_synthesis, normal))
            story.append(Spacer(1, 8))

        # Critères
        story.append(Paragraph("Évaluation détaillée par critère", h2_style))
        for c in result.criteres:
            bar = "█" * int(c.score / 10) + "░" * (10 - int(c.score / 10))
            story.append(Paragraph(
                f"<b>{c.nom}</b> ({c.poids}%) : {c.score:.0f}/100  {bar}",
                normal
            ))
            story.append(Paragraph(f"  → {c.justification}", normal))
            story.append(Spacer(1, 4))

        # Blockers
        if result.criteres_bloquants:
            story.append(Spacer(1, 8))
            story.append(Paragraph("Criteres bloquants", h2_style))
            for b in result.criteres_bloquants:
                story.append(Paragraph(f"⛔ {b}", ParagraphStyle("Blocker", parent=normal, textColor=colors.red)))
            story.append(Spacer(1, 8))

        # Forces / faiblesses
        if result.forces:
            story.append(Paragraph("Points forts", h2_style))
            for f in result.forces:
                story.append(Paragraph(f"✓ {f}", normal))
        if result.faiblesses:
            story.append(Paragraph("Points faibles", h2_style))
            for f in result.faiblesses:
                story.append(Paragraph(f"✗ {f}", normal))

        # Recommandations
        story.append(Spacer(1, 8))
        story.append(Paragraph("Recommandations", h2_style))
        for r in result.recommandations:
            story.append(Paragraph(f"• {r}", normal))

        # Valeur ajoutée IA
        valeur = ai.get("valeur_ajoutee")
        if valeur:
            story.append(Spacer(1, 8))
            story.append(Paragraph("Notre valeur ajoutée", h2_style))
            items = valeur if isinstance(valeur, list) else [valeur]
            for item in items:
                story.append(Paragraph(f"✓ {item}", normal))

        # Conclusion IA
        conclusion = ai.get("conclusion")
        if conclusion:
            story.append(Spacer(1, 8))
            story.append(Paragraph("Conclusion", h2_style))
            story.append(Paragraph(conclusion, normal))

        doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=2*cm, leftMargin=2*cm)
        doc.build(story)
        return path
